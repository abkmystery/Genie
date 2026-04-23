from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from app.providers.ocr_provider import MetadataOCRProvider
from app.providers.source_parsers import (
    CsvSourceParser,
    DocxSourceParser,
    ImageSourceParser,
    PdfSourceParser,
    TextSourceParser,
    XlsxSourceParser,
)


def test_text_and_markdown_parser(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text("# Heading\nFastAPI and desktop workflows", encoding="utf-8")

    chunks = TextSourceParser().parse(path)
    assert chunks
    assert "FastAPI" in chunks[0].text


def test_csv_parser(tmp_path: Path):
    path = tmp_path / "budget.csv"
    path.write_text("category,amount\ntravel,1200\nsoftware,450\n", encoding="utf-8")

    chunks = CsvSourceParser().parse(path)
    assert len(chunks) == 2
    assert chunks[0].locator == "row-range=1-1"


def test_xlsx_parser(tmp_path: Path):
    path = tmp_path / "budget.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["category", "amount"])
    sheet.append(["travel", 1200])
    workbook.save(path)

    chunks = XlsxSourceParser().parse(path)
    assert any("travel" in chunk.text for chunk in chunks)


def test_docx_parser(tmp_path: Path):
    path = tmp_path / "handbook.docx"
    document = Document()
    document.add_paragraph("Benefits begin after onboarding.")
    document.save(path)

    chunks = DocxSourceParser().parse(path)
    assert chunks
    assert "Benefits begin" in chunks[0].text


def test_pdf_parser(tmp_path: Path):
    path = tmp_path / "guide.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Genie PDF guide")
    pdf.save()

    chunks = PdfSourceParser().parse(path)
    assert chunks
    assert "Genie PDF guide" in chunks[0].text


def test_image_parser_uses_ocr_fallback(tmp_path: Path):
    from PIL import Image

    path = tmp_path / "diagram.png"
    Image.new("RGB", (200, 120), color=(250, 250, 250)).save(path)

    chunks = ImageSourceParser(MetadataOCRProvider()).parse(path)
    assert chunks
    assert "diagram.png" in chunks[0].text

